"""
Excel report exporter for CO-PO mapping results.

Generates:
- Mapping matrix sheet (CO rows × PO/PSO columns)
- Detailed mappings sheet
- Evaluation results sheet
- Styled with color-coded mapping levels
"""

import logging
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Color scheme for mapping levels
LEVEL_FILLS = {
    0: PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"),  # Grey
    1: PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid"),  # Light Blue
    2: PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"),  # Light Yellow
    3: PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),  # Light Green
}

HEADER_FILL = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CELL_FONT = Font(size=11)
CELL_ALIGNMENT = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)


def _level_display(level: int) -> str:
    """Convert level to display value."""
    return "-" if level == 0 else str(level)


def export_mapping_report(
    file_path: str,
    subject_name: str,
    co_codes: List[str],
    po_codes: List[str],
    pso_codes: List[str],
    ai_mappings: Dict[Tuple[str, str], Dict],
    reference_mappings: Optional[Dict[Tuple[str, str], int]] = None,
    evaluation: Optional[Dict] = None,
):
    """
    Export a complete mapping report to Excel.

    Args:
        file_path: Output file path.
        subject_name: Course/subject name.
        co_codes: List of CO codes.
        po_codes: List of PO codes.
        pso_codes: List of PSO codes.
        ai_mappings: Dict mapping (co_code, po_pso_code) -> {level, confidence, reason, ...}
        reference_mappings: Optional faculty reference mappings.
        evaluation: Optional evaluation results dict.
    """
    wb = openpyxl.Workbook()

    # ---- Sheet 1: Mapping Matrix ----
    ws_matrix = wb.active
    ws_matrix.title = "CO-PO Mapping Matrix"

    all_po_pso = po_codes + pso_codes

    # Title row
    ws_matrix.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(all_po_pso) + 1)
    title_cell = ws_matrix.cell(row=1, column=1, value=f"CO-PO/PSO Mapping Matrix — {subject_name}")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Headers
    ws_matrix.cell(row=3, column=1, value="CO \\ PO-PSO").font = HEADER_FONT
    ws_matrix.cell(row=3, column=1).fill = HEADER_FILL
    ws_matrix.cell(row=3, column=1).border = THIN_BORDER

    for j, code in enumerate(all_po_pso, start=2):
        cell = ws_matrix.cell(row=3, column=j, value=code)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CELL_ALIGNMENT
        cell.border = THIN_BORDER
        ws_matrix.column_dimensions[get_column_letter(j)].width = 8

    ws_matrix.column_dimensions['A'].width = 15

    # Data rows
    for i, co_code in enumerate(co_codes, start=4):
        co_cell = ws_matrix.cell(row=i, column=1, value=co_code)
        co_cell.font = Font(bold=True)
        co_cell.border = THIN_BORDER

        for j, po_pso_code in enumerate(all_po_pso, start=2):
            key = (co_code, po_pso_code)
            mapping_data = ai_mappings.get(key, {})
            level = mapping_data.get("final_level", mapping_data.get("level", 0))

            cell = ws_matrix.cell(row=i, column=j, value=_level_display(level))
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGNMENT
            cell.fill = LEVEL_FILLS.get(level, LEVEL_FILLS[0])
            cell.border = THIN_BORDER

    # ---- Sheet 2: Detailed Mappings ----
    ws_detail = wb.create_sheet("Detailed Mappings")

    detail_headers = [
        "CO", "PO/PSO", "AI Level", "Confidence",
        "Similarity", "AI Explanation", "Faculty Decision",
        "Final Level", "Faculty Notes"
    ]

    for j, header in enumerate(detail_headers, start=1):
        cell = ws_detail.cell(row=1, column=j, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    ws_detail.column_dimensions['A'].width = 12
    ws_detail.column_dimensions['B'].width = 10
    ws_detail.column_dimensions['C'].width = 10
    ws_detail.column_dimensions['D'].width = 12
    ws_detail.column_dimensions['E'].width = 12
    ws_detail.column_dimensions['F'].width = 60
    ws_detail.column_dimensions['G'].width = 16
    ws_detail.column_dimensions['H'].width = 12
    ws_detail.column_dimensions['I'].width = 30

    row_idx = 2
    for co_code in co_codes:
        for po_pso_code in all_po_pso:
            key = (co_code, po_pso_code)
            data = ai_mappings.get(key, {})

            ws_detail.cell(row=row_idx, column=1, value=co_code).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=2, value=po_pso_code).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=3, value=_level_display(data.get("level", 0))).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=4, value=data.get("confidence", 0)).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=5, value=data.get("similarity", 0)).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=6, value=data.get("reason", "")).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=7, value=data.get("faculty_decision", "pending")).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=8, value=_level_display(data.get("final_level", data.get("level", 0)))).border = THIN_BORDER
            ws_detail.cell(row=row_idx, column=9, value=data.get("faculty_notes", "")).border = THIN_BORDER

            # Color code the AI level cell
            level = data.get("level", 0)
            ws_detail.cell(row=row_idx, column=3).fill = LEVEL_FILLS.get(level, LEVEL_FILLS[0])

            row_idx += 1

    # ---- Sheet 3: Reference Comparison (if available) ----
    if reference_mappings:
        ws_ref = wb.create_sheet("Reference Comparison")

        ref_headers = ["CO", "PO/PSO", "AI Level", "Faculty Level", "Match"]
        for j, header in enumerate(ref_headers, start=1):
            cell = ws_ref.cell(row=1, column=j, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER

        row_idx = 2
        for co_code in co_codes:
            for po_pso_code in all_po_pso:
                key = (co_code, po_pso_code)
                ai_data = ai_mappings.get(key, {})
                ai_level = ai_data.get("level", 0)
                ref_level = reference_mappings.get(key, None)

                if ref_level is not None:
                    match = "✓" if ai_level == ref_level else "✗"
                    match_fill = PatternFill(
                        start_color="C8E6C9" if ai_level == ref_level else "FFCDD2",
                        end_color="C8E6C9" if ai_level == ref_level else "FFCDD2",
                        fill_type="solid"
                    )

                    ws_ref.cell(row=row_idx, column=1, value=co_code).border = THIN_BORDER
                    ws_ref.cell(row=row_idx, column=2, value=po_pso_code).border = THIN_BORDER
                    ws_ref.cell(row=row_idx, column=3, value=_level_display(ai_level)).border = THIN_BORDER
                    ws_ref.cell(row=row_idx, column=4, value=_level_display(ref_level)).border = THIN_BORDER
                    match_cell = ws_ref.cell(row=row_idx, column=5, value=match)
                    match_cell.border = THIN_BORDER
                    match_cell.fill = match_fill
                    match_cell.alignment = CELL_ALIGNMENT

                    row_idx += 1

    # ---- Sheet 4: Evaluation Results (if available) ----
    if evaluation:
        ws_eval = wb.create_sheet("Evaluation Results")

        eval_items = [
            ("Metric", "Value"),
            ("Accuracy", evaluation.get("accuracy", "N/A")),
            ("Precision (Macro)", evaluation.get("precision_macro", "N/A")),
            ("Recall (Macro)", evaluation.get("recall_macro", "N/A")),
            ("F1 Score (Macro)", evaluation.get("f1_macro", "N/A")),
            ("Agreement %", evaluation.get("agreement_percentage", "N/A")),
            ("Total Pairs", evaluation.get("total_pairs", "N/A")),
            ("Matched Pairs", evaluation.get("matched_pairs", "N/A")),
        ]

        for i, (metric, value) in enumerate(eval_items, start=1):
            cell_m = ws_eval.cell(row=i, column=1, value=metric)
            cell_v = ws_eval.cell(row=i, column=2, value=value)
            cell_m.border = THIN_BORDER
            cell_v.border = THIN_BORDER
            if i == 1:
                cell_m.font = HEADER_FONT
                cell_m.fill = HEADER_FILL
                cell_v.font = HEADER_FONT
                cell_v.fill = HEADER_FILL

        ws_eval.column_dimensions['A'].width = 25
        ws_eval.column_dimensions['B'].width = 20

    # Save
    wb.save(file_path)
    logger.info(f"Exported mapping report to: {file_path}")
