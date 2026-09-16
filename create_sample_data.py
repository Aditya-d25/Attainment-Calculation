"""
Generate a sample reference Excel workbook for testing.

Contains CO-PO mapping data similar to what a faculty member would create.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def create_sample_reference():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO-DAttainment"

    # ---- Headers ----
    headers = ["CO", "Description", "PO1", "PO2", "PO3", "PO4", "PO5",
               "PO6", "PO7", "PO8", "PO9", "PO10", "PO11", "PSO1", "PSO2", "PSO3"]

    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # ---- CO Data ----
    # Data Structures course COs with reference mappings
    cos = [
        {
            "code": "CO206.1",
            "desc": "Apply and implement the principle sorting and searching algorithms on the given data using the C language.",
            "mappings": [2, 1, 1, 1, 1, "-", 1, 1, "-", "-", 1, "-", 1, "-"]
        },
        {
            "code": "CO206.2",
            "desc": "Develop applications of stack and queue using arrays.",
            "mappings": [2, 2, 2, 2, 1, "-", "-", "-", 1, "-", 1, "-", 1, "-"]
        },
        {
            "code": "CO206.3",
            "desc": "Implement and demonstrate the applicability of a Linked List.",
            "mappings": [2, 2, 2, 2, 1, "-", "-", "-", 1, "-", 1, "-", 1, "-"]
        },
        {
            "code": "CO206.4",
            "desc": "Use trees and graphs for solving real-world problems.",
            "mappings": [2, 2, 2, 1, 1, "-", "-", "-", "-", "-", 1, 1, 1, "-"]
        },
        {
            "code": "CO206.5",
            "desc": "Analyze the efficiency of algorithms using time and space complexity measures.",
            "mappings": [3, 2, 1, 2, 1, "-", "-", "-", "-", "-", 1, 1, 1, "-"]
        },
        {
            "code": "CO206.6",
            "desc": "Design and implement hashing techniques for efficient data retrieval.",
            "mappings": [2, 1, 3, 1, 2, "-", "-", "-", "-", "-", 1, "-", 2, "-"]
        },
    ]

    for i, co in enumerate(cos, 2):
        ws.cell(row=i, column=1, value=co["code"]).border = thin_border
        ws.cell(row=i, column=2, value=co["desc"]).border = thin_border
        ws.column_dimensions['B'].width = 60

        for j, val in enumerate(co["mappings"], 3):
            cell = ws.cell(row=i, column=j)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if val == "-":
                cell.value = "-"
            else:
                cell.value = val

    ws.column_dimensions['A'].width = 12

    # Save
    output_path = "sample_data/reference_mapping.xlsx"
    wb.save(output_path)
    print(f"Created sample reference: {output_path}")
    return output_path


if __name__ == "__main__":
    import os
    os.makedirs("sample_data", exist_ok=True)
    create_sample_reference()
